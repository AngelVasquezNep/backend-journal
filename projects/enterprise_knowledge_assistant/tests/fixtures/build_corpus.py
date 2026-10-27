"""Build the versioned two-company document corpus without external network access.

Run with the Knowledge environment. Output is deterministic text and generated local
documents under corpus/files; only the manifest is the evaluation authority.
"""

import argparse
import json
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter


def native_pdf(text: str) -> bytes:
    """Encode ASCII fixture text as a one-page native PDF with a stable page locator."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 40 740 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    result = b"%PDF-1.4\n"
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(result))
        result += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    start = len(result)
    result += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets[1:]:
        result += f"{offset:010d} 00000 n \n".encode()
    return (
        result
        + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    )


def text_image(text: str) -> Image.Image:
    """Render large printed text on a white image for the real OCR acceptance cases."""
    image = Image.new("RGB", (1800, 400), "white")
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 36)
    except OSError:
        font = ImageFont.load_default(size=36)
    words = text.split()
    lines = [" ".join(words[start : start + 9]) for start in range(0, len(words), 9)]
    ImageDraw.Draw(image).multiline_text(
        (40, 40), "\n".join(lines), fill="black", font=font, spacing=16
    )
    return image


def build(destination: Path) -> None:
    """Write supported, conflicting, malicious and invalid files from the frozen manifest."""
    manifest = json.loads((Path(__file__).parent / "corpus/manifest.json").read_text())
    destination.mkdir(parents=True, exist_ok=True)
    for entry in manifest["documents"]:
        target = destination / entry["file"]
        target.parent.mkdir(parents=True, exist_ok=True)
        text = entry["text"]
        match entry["format"]:
            case "txt":
                target.write_text(text, encoding="utf-8")
            case "docx":
                document = Document()
                document.add_heading(entry["title"], level=1)
                table = document.add_table(rows=1, cols=1)
                table.cell(0, 0).text = text
                document.save(str(target))
            case "pdf":
                target.write_bytes(native_pdf(text))
            case "scanned_pdf":
                text_image(text).save(target, "PDF", resolution=150)
            case "png" | "jpeg":
                text_image(text).save(target)
            case _:
                raise ValueError("Unsupported manifest format")
    (destination / "invalid").mkdir(exist_ok=True)
    (destination / "invalid/empty.txt").write_bytes(b"")
    (destination / "invalid/one-byte.txt").write_bytes(b"x")
    (destination / "invalid/corrupt.pdf").write_bytes(b"%PDF-not-a-document")
    (destination / "invalid/mismatch.png").write_bytes(b"plain text, not a PNG")
    (destination / "invalid/not-utf8.txt").write_bytes(b"\xff\xfe\x00")
    writer = PdfWriter()
    source = next(entry for entry in manifest["documents"] if entry["format"] == "pdf")
    writer.append(PdfReader(destination / source["file"]))
    writer.encrypt("fixture-password-not-an-account-credential")
    with (destination / "invalid/protected.pdf").open("wb") as stream:
        writer.write(stream)
    (destination / "invalid/exactly-20-mib.txt").write_bytes(b"x" * 20_971_520)
    (destination / "invalid/over-20-mib.txt").write_bytes(b"x" * 20_971_521)


def main() -> None:
    """Parse the local output directory and build fixtures; failures propagate to the operator."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output", type=Path, default=Path(__file__).parent / "corpus/files"
    )
    build(parser.parse_args().output)


if __name__ == "__main__":
    main()
